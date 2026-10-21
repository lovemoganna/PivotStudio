"""Prepare an unchanged public workbook and independent streaming oracle."""
import hashlib
import json
import urllib.request
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

folder = Path('artifacts/business-data')
folder.mkdir(parents=True, exist_ok=True)
archive = folder / 'retail-source.zip'
url = 'https://archive.ics.uci.edu/static/public/352/online+retail.zip'
if not archive.exists():
    with urllib.request.urlopen(url, timeout=60) as response:
        archive.write_bytes(response.read())
digest = hashlib.sha256(archive.read_bytes()).hexdigest()
if digest != 'f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b':
    raise ValueError('UCI source differs from the recorded baseline; review it before comparing results')
file = folder / 'retail-full.xlsx'
with zipfile.ZipFile(archive) as package:
    file.write_bytes(package.read(next(name for name in package.namelist() if name.lower().endswith('.xlsx'))))
book = load_workbook(file, read_only=True, data_only=True)
records = book.active.iter_rows(values_only=True)
headers = next(records)
groups = defaultdict(lambda: {'count': 0, 'quantity': 0})
missing = Counter()
quantities = Counter()
count = quantity = 0
first = None
for values in records:
    row = dict(zip(headers, values))
    count += 1
    qty = row.get('Quantity') or 0
    quantity += qty
    quantities[qty] += 1
    group = groups[row.get('Country') or '(空)']
    group['count'] += 1
    group['quantity'] += qty
    missing.update(key for key, value in row.items() if value is None)
    if first is None:
        first = {key: value.isoformat() if hasattr(value, 'isoformat') else value for key, value in row.items()}
book.close()
result = {'sourceURL': url, 'sourceSHA256': digest, 'file': str(file), 'fileSHA256': hashlib.sha256(file.read_bytes()).hexdigest(), 'rows': count, 'quantity': quantity, 'groups': groups, 'quantityFrequency': dict(quantities), 'missing': dict(missing), 'first': first}
(folder / 'full-oracle.json').write_text(json.dumps(result, indent=2))
print(json.dumps({'rows': count, 'quantity': quantity, 'countries': len(groups), 'missing': missing}), flush=True)
