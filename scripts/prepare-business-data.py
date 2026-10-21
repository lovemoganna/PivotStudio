"""Download public UCI business data; never send product/user data anywhere."""
import csv
import hashlib
import io
import json
from pathlib import Path
import urllib.request
import zipfile
from decimal import Decimal
from itertools import islice
from openpyxl import load_workbook, Workbook

OUT = Path('artifacts/business-data')
OUT.mkdir(parents=True, exist_ok=True)
SOURCES = [
    (352, 'online+retail', 'retail'),
    (222, 'bank+marketing', 'bank'),
    (468, 'online+shoppers+purchasing+intention+dataset', 'shoppers'),
]
manifest = []
for ident, slug, name in SOURCES:
    url = f'https://archive.ics.uci.edu/static/public/{ident}/{slug}.zip'
    archive = OUT / f'{name}-source.zip'
    if not archive.exists():
        with urllib.request.urlopen(url, timeout=45) as response:
            archive.write_bytes(response.read())
    with zipfile.ZipFile(archive) as package:
        if name == 'retail':
            raw = package.read(next(p for p in package.namelist() if p.lower().endswith('.xlsx')))
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            # Keep the first 2,000 source records without changing values or headers.
            source_rows = list(islice(book.active.iter_rows(values_only=True), 2001))
            target = Workbook(); sheet = target.active
            for row in source_rows:
                sheet.append(row)
            file = OUT / 'retail-first-2000.xlsx'; target.save(file)
            headers = source_rows[0]
            rows = [dict(zip(headers, row)) for row in source_rows[1:]]
            group, measure, agg, selected = 'Country', 'Quantity', 'sum', 'United Kingdom'
            extra = {'subset': 'First 2000 records in source worksheet order, unchanged cell values', 'priceCurrency': 'GBP'}
            book.close()
        else:
            if name == 'bank':
                nested = zipfile.ZipFile(io.BytesIO(package.read('bank.zip')))
                raw = nested.read('bank-full.csv'); delimiter = ';'; file = OUT / 'bank-full.csv'
                group, measure, agg, selected = 'job', 'balance', 'sum', 'management'
            else:
                raw = package.read(next(p for p in package.namelist() if p.endswith('.csv')))
                delimiter = ','; file = OUT / 'online-shoppers.csv'
                group, measure, agg, selected = 'Month', 'BounceRates', 'avg', 'May'
            file.write_bytes(raw)
            rows = list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig')), delimiter=delimiter))
            extra = {'subset': 'Complete published CSV'}
        matching = [row for row in rows if str(row[group]) == selected]
        values = [Decimal(str(row[measure])) for row in matching if row[measure] is not None and str(row[measure]) != '']
        result = sum(values) if agg == 'sum' else sum(values) / len(values)
        item = {'name': name, 'url': url, 'sourcePage': f'https://archive.ics.uci.edu/dataset/{ident}/{slug}', 'sourceSHA256': hashlib.sha256(archive.read_bytes()).hexdigest(), 'file': str(file), 'fileSHA256': hashlib.sha256(file.read_bytes()).hexdigest(), 'rows': len(rows), 'group': group, 'measure': measure, 'agg': agg, 'selected': selected, 'filteredRows': len(matching), 'expected': str(result), **extra}
        manifest.append(item)
        print(json.dumps(item), flush=True)
(OUT / 'manifest.json').write_text(json.dumps(manifest, indent=2))
